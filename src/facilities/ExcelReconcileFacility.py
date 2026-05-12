from datetime import datetime
from pathlib import Path
from openpyxl import Workbook
from openpyxl import load_workbook
from src.structures.StructReconcileRecord import StructReconcileRecord

class ExcelReconcileService:

    def process_files(self,input_folder,
        output_folder,
        event_handler=None
    ):
        # Extract file names and their count
        input_path = Path(input_folder)

        if not output_folder:
            current_datetime = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_folder_name = (f'output_{current_datetime}')
            output_folder_path = (input_path/ output_folder_name)
            try:
                output_folder_path.mkdir(exist_ok=False)
                output_folder=output_folder_path
            except Exception as ex:
                if event_handler:

                    event_handler(
                            {
                                'event': 'folder_creation_failed',
                                'message': (
                                    "Output folder can't "
                                    "be created"
                                ),
                                'error': str(ex)
                            }
                        )                    
                    raise Exception(f'Failed to create output folder: ' f'{str(ex)}')
        excel_files = list(
            input_path.glob('*.xls*')
        )
        total_files = len(excel_files)

        # Gathering data
        consolidated_records = []

        for index, file in enumerate(excel_files):

            record = self.read_file(file)

            consolidated_records.append(
                record
            )
            if event_handler:
                event_handler(
                    {
                        'event': 'progress_changed',
                        'current': index + 1,
                        'total': total_files,
                        'file_name': file.name,
                        'status': record.status
                    }
                )

        output_file = self.write_output_file(
            consolidated_records,
            input_folder,
            output_folder
        )

        if event_handler:
            event_handler(
                {
                    'event': 'completed',
                    'output_file': str(output_file),
                    'total_records': len(consolidated_records)
                }
            )
        return output_file

    def read_file(
        self,
        file_path
    ):

        record = StructReconcileRecord()
        try:
            workbook = load_workbook(
                file_path,
                data_only=True
            )

            if 'Sheet1' not in workbook.sheetnames:
                raise Exception('Sheet1 not found')
            sheet = workbook['Sheet1']
            record.file_name = file_path.name
            record.m13 = sheet['M13'].value
            record.g14 = sheet['G14'].value
            record.status = 'SUCCESS'
            record.error_message = ''

        except Exception as ex:
            record.file_name = file_path.name
            record.m13 = ''
            record.g14 = ''
            record.status = 'FAILED'
            record.error_message = str(ex)
        return record

    def write_output_file(self,records,
        input_folder,
        output_folder
    ):
        # Create new workbook file
        workbook = Workbook()

        sheet = workbook.active
        sheet.title = 'Output Data'

        headers = [
            'File Name',
            'M13',
            'G14',
            'Status',
            'Error Message'
        ]
        sheet.append(headers)
        for record in records:

            sheet.append([
                record.file_name,
                record.m13,
                record.g14,
                record.status,
                record.error_message
            ])

        input_folder_name = Path(input_folder).name
        current_datetime = datetime.now().strftime(
        '%Y%m%d_%H%M%S'
        )

        output_file_name = (
            f'out_'
            f'{input_folder_name}_'
            f'{current_datetime}.xlsx'
        )

        output_file = (
                Path(output_folder)
                / output_file_name
            )
        workbook.save(output_file)
        return output_file